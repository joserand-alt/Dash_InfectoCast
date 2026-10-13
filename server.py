import http.server
import socketserver
import json
import subprocess
import os
import sys
import openpyxl

PORT = 8085
EXCEL_PATH = r'C:\Users\DELL\Desktop\Acompanhamento de acessos\BD\CURSOS.xlsx'

class APIHandler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, Authorization')
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def do_GET(self):
        if self.path == '/api/status':
            self.send_response(200)
            self.send_header('Content-type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps({"status": "running", "port": PORT}).encode('utf-8'))
            return
        super().do_GET()

    def do_POST(self):
        if self.path == '/api/alocar-aula':
            content_length = int(self.headers['Content-Length'])
            post_data = self.rfile.read(content_length)
            data = json.loads(post_data.decode('utf-8'))
            
            curso = data.get('curso')
            aula = data.get('aula')
            modulo_nome = data.get('modulo')
            
            try:
                # Carregar planilha usando openpyxl para manter formatação
                wb = openpyxl.load_workbook(EXCEL_PATH)
                ws_mods = wb.worksheets[1] # Aba "Módulos"
                ws_aulas = wb.worksheets[2] # Aba "Aulas"
                
                # 1. Achar ou criar o Módulo
                modulo_id = None
                max_mod_id = 0
                
                # Procura o módulo pelo nome
                for row in ws_mods.iter_rows(min_row=2, values_only=False):
                    m_id_val = row[1].value
                    m_nome_val = row[3].value
                    
                    if isinstance(m_id_val, (int, float)):
                        max_mod_id = max(max_mod_id, int(m_id_val))
                    elif isinstance(m_id_val, str) and m_id_val.isdigit():
                        max_mod_id = max(max_mod_id, int(m_id_val))
                        
                    if m_nome_val and str(m_nome_val).strip().lower() == modulo_nome.lower():
                        modulo_id = m_id_val
                        break
                        
                if not modulo_id:
                    modulo_id = max_mod_id + 1
                    # Append new module
                    # Columns: ['CURSO', 'ID MÓDULO', 'Ordem MÓDULO', 'Nome módulo', 'Descrição', 'Aulas', 'Datas Turmas', 'Ativo']
                    ws_mods.append([curso, modulo_id, "", modulo_nome, "", "", "", ""])
                    
                # 2. Adicionar a Aula
                # Precisamos de um ID para a aula também
                max_aula_id = 0
                for row in ws_aulas.iter_rows(min_row=2, values_only=True):
                    a_id_val = row[2]
                    if isinstance(a_id_val, (int, float)):
                        max_aula_id = max(max_aula_id, int(a_id_val))
                    elif isinstance(a_id_val, str) and a_id_val.isdigit():
                        max_aula_id = max(max_aula_id, int(a_id_val))
                        
                nova_aula_id = max_aula_id + 1
                
                # Columns: ['CURSO', 'MÓDULO' (ID), 'ID' (Aula ID), 'Ordem', 'Nome da aula', 'Vídeo', 'Materiais', 'Atividades', 'Ativo']
                ws_aulas.append([curso, modulo_id, nova_aula_id, "", aula, "", "", "", ""])
                
                # Salvar a planilha
                wb.save(EXCEL_PATH)
                
                # Regerar o dashboard
                subprocess.run(['python', 'gerador.py'], check=True)
                
                self.send_response(200)
                self.send_header('Content-type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps({"status": "ok"}).encode('utf-8'))
                
            except Exception as e:
                print(f"Erro: {e}")
                self.send_response(500)
                self.send_header('Content-type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps({"error": str(e)}).encode('utf-8'))
        elif self.path == '/api/atualizar':
            try:
                base_dir = os.path.dirname(os.path.abspath(__file__))
                print("[SERVER] Iniciando atualização manual completa de todas as APIs...")
                
                # 1. Sincronizar matrículas e disparar tags no RD Station
                try:
                    print("[SERVER] 1/3 Sincronizando tags e matrículas no RD Station...")
                    subprocess.run([sys.executable, os.path.join(base_dir, 'sync_matriculas_rd.py')], cwd=base_dir, check=False)
                except Exception as e_rd:
                    print(f"[SERVER] Aviso RD: {e_rd}")

                # 2. Consultar todas as APIs (Cativa, Vindi, Asaas, Academy, RD) e regerar dashboard
                print("[SERVER] 2/3 Consultando APIs e gerando HTML...")
                subprocess.run([sys.executable, os.path.join(base_dir, 'gerador.py')], cwd=base_dir, check=True)

                # 3. Garantir cópia para index.html
                gen_path = os.path.join(base_dir, 'dashboard_gerado.html')
                idx_path = os.path.join(base_dir, 'index.html')
                if os.path.exists(gen_path):
                    with open(gen_path, 'rb') as src, open(idx_path, 'wb') as dst:
                        dst.write(src.read())

                # 4. Salvar e enviar para o GitHub Pages para atualizar online
                try:
                    print("[SERVER] 3/3 Enviando para o GitHub Pages...")
                    desktop_dir = r'C:\Users\DELL\Desktop\Dash_InfectoCast'
                    if os.path.exists(desktop_dir) and os.path.abspath(base_dir).lower() != os.path.abspath(desktop_dir).lower():
                        import shutil
                        for fname in ['index.html', 'dashboard_gerado.html', 'template.html', 'asaas_cache.json', 'vindi_cache.json', 'cativa_cache.json']:
                            sf = os.path.join(base_dir, fname)
                            df = os.path.join(desktop_dir, fname)
                            if os.path.exists(sf):
                                shutil.copy2(sf, df)

                    subprocess.run(['git', 'add', 'index.html', 'dashboard_gerado.html', 'template.html', '*.json'], cwd=base_dir, check=False)
                    subprocess.run(['git', 'commit', '-m', 'Manual-Sync: Dados atualizados via botão do dashboard [skip ci]'], cwd=base_dir, check=False)
                    push_res = subprocess.run(['git', 'push', 'origin', 'main'], cwd=base_dir, check=False)
                    if push_res.returncode != 0:
                        subprocess.run(['git', 'pull', '--no-rebase', '-X', 'theirs', 'origin', 'main'], cwd=base_dir, check=False)
                        subprocess.run(['git', 'push', 'origin', 'main'], cwd=base_dir, check=False)
                    print("[SERVER] Sucesso: Dados enviados ao GitHub Pages!")
                except Exception as e_git:
                    print(f"[SERVER] Aviso Git push: {e_git}")

                self.send_response(200)
                self.send_header('Content-type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps({
                    "status": "ok",
                    "message": "Todas as APIs foram consultadas, os dados foram salvos e o painel foi atualizado com sucesso!"
                }).encode('utf-8'))
            except Exception as e:
                print(f"[SERVER] Erro na atualização: {e}")
                self.send_response(500)
                self.send_header('Content-type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps({"error": str(e)}).encode('utf-8'))
        else:
            self.send_response(404)
            self.end_headers()

# Prevent address already in use by killing old server if possible
# (Handled externally via manage_task)
socketserver.TCPServer.allow_reuse_address = True
with socketserver.TCPServer(("", PORT), APIHandler) as httpd:
    print(f"Servidor backend rodando na porta {PORT}...")
    httpd.serve_forever()
